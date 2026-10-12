"""
Export a generated timetable to PDF and Excel. Both formats render the same
data: one grid per class (rows = periods, columns = days), so what you see
in the frontend viewer's class view matches what you download.
"""
from __future__ import annotations

import io

from fastapi import APIRouter, Depends, HTTPException, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import select
from sqlalchemy.orm import Session, joinedload

from app.auth import get_current_tenant
from app.db.session import get_db
from app.models.day import Day
from app.models.period import Period
from app.models.school_class import SchoolClass
from app.models.tenant import Tenant
from app.models.timetable import Timetable, TimetableEntry

router = APIRouter(prefix="/api/timetables", tags=["export"])


def _grid_data(
    db: Session, timetable_id: int, tenant_id: int
) -> tuple[list[SchoolClass], list[Day], list[Period], dict]:
    classes = list(
        db.scalars(
            select(SchoolClass).where(SchoolClass.tenant_id == tenant_id).order_by(SchoolClass.grade, SchoolClass.section)
        ).all()
    )
    days = list(db.scalars(select(Day).where(Day.tenant_id == tenant_id).order_by(Day.order_index)).all())
    periods = list(
        db.scalars(
            select(Period).where(Period.is_break.is_(False), Period.tenant_id == tenant_id).order_by(Period.period_number)
        ).all()
    )
    entries = list(
        db.scalars(
            select(TimetableEntry)
            .where(TimetableEntry.timetable_id == timetable_id)
            .options(
                joinedload(TimetableEntry.subject),
                joinedload(TimetableEntry.teacher),
                joinedload(TimetableEntry.room),
            )
        ).all()
    )
    # (class_id, day_id, period_id) -> entry
    by_slot = {(e.class_id, e.day_id, e.period_id): e for e in entries}
    return classes, days, periods, by_slot


def _cell_text(entry: TimetableEntry | None) -> str:
    if entry is None:
        return "-"
    return f"{entry.subject.name}\n{entry.teacher.name}\n{entry.room.name}"


@router.get("/{timetable_id}/export/excel")
def export_excel(
    timetable_id: int, db: Session = Depends(get_db), tenant: Tenant = Depends(get_current_tenant)
) -> StreamingResponse:
    timetable = db.get(Timetable, timetable_id)
    if timetable is None or timetable.tenant_id != tenant.id:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=f"Timetable {timetable_id} not found")

    classes, days, periods, by_slot = _grid_data(db, timetable_id, tenant.id)

    wb = Workbook()
    wb.remove(wb.active)
    header_fill = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    wrap = Alignment(wrap_text=True, vertical="center", horizontal="center")

    for school_class in classes:
        sheet_name = school_class.display_name[:31]  # Excel sheet name limit
        ws = wb.create_sheet(title=sheet_name)

        ws.cell(row=1, column=1, value="Period / Day").font = header_font
        ws.cell(row=1, column=1).fill = header_fill
        for col, day in enumerate(days, start=2):
            cell = ws.cell(row=1, column=col, value=day.name)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = wrap

        for row, period in enumerate(periods, start=2):
            label = f"P{period.period_number}\n{period.start_time.strftime('%H:%M')}-{period.end_time.strftime('%H:%M')}"
            label_cell = ws.cell(row=row, column=1, value=label)
            label_cell.alignment = wrap
            label_cell.font = Font(bold=True)
            for col, day in enumerate(days, start=2):
                entry = by_slot.get((school_class.id, day.id, period.id))
                cell = ws.cell(row=row, column=col, value=_cell_text(entry))
                cell.alignment = wrap

        ws.column_dimensions["A"].width = 16
        for col in range(2, len(days) + 2):
            ws.column_dimensions[get_column_letter(col)].width = 22
        for row in range(2, len(periods) + 2):
            ws.row_dimensions[row].height = 48

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    filename = f"timetable_{timetable_id}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/{timetable_id}/export/pdf")
def export_pdf(
    timetable_id: int, db: Session = Depends(get_db), tenant: Tenant = Depends(get_current_tenant)
) -> StreamingResponse:
    timetable = db.get(Timetable, timetable_id)
    if timetable is None or timetable.tenant_id != tenant.id:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=f"Timetable {timetable_id} not found")

    classes, days, periods, by_slot = _grid_data(db, timetable_id, tenant.id)

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(letter), topMargin=24, bottomMargin=24)
    styles = getSampleStyleSheet()
    story = []

    for i, school_class in enumerate(classes):
        story.append(Paragraph(f"{school_class.display_name} - {timetable.name}", styles["Title"]))
        story.append(Spacer(1, 12))

        header = ["Period"] + [d.name for d in days]
        rows = [header]
        for period in periods:
            row = [f"P{period.period_number}\n{period.start_time.strftime('%H:%M')}"]
            for day in days:
                entry = by_slot.get((school_class.id, day.id, period.id))
                row.append(_cell_text(entry))
            rows.append(row)

        table = Table(rows, repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E3A5F")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTSIZE", (0, 0), (-1, -1), 7),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F5F8")]),
                ]
            )
        )
        story.append(table)
        if i < len(classes) - 1:
            story.append(PageBreak())

    doc.build(story)
    buffer.seek(0)
    filename = f"timetable_{timetable_id}.pdf"
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
